"""Candidate Crossref references to pinned historical NHMRC application cells.
SPDX-License-Identifier: Apache-2.0
"""
from __future__ import annotations

import copy
import hashlib
from io import BytesIO
import json
from pathlib import Path
import posixpath
import re
import xml.etree.ElementTree as ET
from zipfile import ZipFile

import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string

if __package__:
    from scripts.award_identifier_crosswalks import authority_contract, exact_text, governing_envelope, identifier
    from scripts.source_relationship_evidence import fingerprint
    from scripts.plan_publication_stub_restore import strict_json_loads
    from scripts.crossref_award_crosswalks import _count, _doi, _complete
else:
    from award_identifier_crosswalks import authority_contract, exact_text, governing_envelope, identifier
    from source_relationship_evidence import fingerprint
    from plan_publication_stub_restore import strict_json_loads
    from crossref_award_crosswalks import _count, _doi, _complete

NS={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
    'r':'http://schemas.openxmlformats.org/package/2006/relationships'}


def _sha(value):
    if not isinstance(value,str) or not re.fullmatch('[0-9a-f]{64}',value):raise ValueError('Independent SHA256 pin required')
    return value


def exact_identifier(value, *, cell_type='s', number_format='General', numeric_policy='reject_numeric'):
    """Conversion of an explicit numeric cell requires the independently pinned policy.

    No inferred leading zeros, integer-valued float conversion, formula cache or
    date/style coercion is admitted as identity.
    """
    if cell_type in ('s','inlineStr') and isinstance(value,str):
        try:return identifier(value)
        except ValueError:return None
    if (numeric_policy=='positive_integer_decimal' and cell_type=='n' and type(value) is int
            and value>0 and value<=2**53 and number_format in ('General','0')):
        return str(value)
    return None


def _safe_xml_declarations(raw):
    # Also catches common UTF-16/32 encodings before any XML/library parser.
    # This is a narrow declaration check, not a general hostile-document sandbox.
    if re.search(br'<!\s*(?:DOCTYPE|ENTITY)\b',raw.replace(b'\x00',b''),re.I):
        raise ValueError('Unsafe XML doctype/entity declarations refused')


def _physical_bounds(xml):
    for row in xml.findall('s:sheetData/s:row',NS):
        text=row.get('r','')
        if not re.fullmatch('[1-9][0-9]*',text) or int(text)>1048576:
            raise ValueError('Worksheet row outside Excel physical bounds')
        for cell in row.findall('s:c',NS):
            match=re.fullmatch(r'([A-Z]+)([1-9][0-9]*)',cell.get('r',''))
            if (not match or int(match[2])!=int(text) or int(match[2])>1048576
                    or len(match[1])>3 or column_index_from_string(match[1])>16384):
                raise ValueError('Worksheet cell outside Excel physical bounds')


def inspect_workbook(data, *, sheet='GRANTS DATA', header_row=1, id_column='APP ID', numeric_policy='reject_numeric'):
    """Inspect all physical rows independent of declared dimensions; never evaluate formulas."""
    if not isinstance(data,bytes) or type(header_row) is not int or header_row<=0 or header_row>1048576:
        raise ValueError('Pinned workbook bytes and explicit header row required')
    if sheet!='GRANTS DATA' or id_column!='APP ID' or numeric_policy not in ('reject_numeric','positive_integer_decimal'):
        raise ValueError('Exact historical application extraction scope required')
    with ZipFile(BytesIO(data)) as archive:
        names=archive.namelist()
        if len(names)!=len(set(names)) or archive.testzip() is not None:raise ValueError('Workbook ZIP integrity or duplicate entry failure')
        worksheet_parts=[]
        for name in names:
            raw_part=archive.read(name)
            _safe_xml_declarations(raw_part)
            if name.startswith('xl/worksheets/') and name.endswith('.xml'):worksheet_parts.append(raw_part)
        for raw_part in worksheet_parts:_physical_bounds(ET.fromstring(raw_part))
        workbook=ET.fromstring(archive.read('xl/workbook.xml'))
        sheets=workbook.findall('s:sheets/s:sheet',NS)
        labels=[s.get('name') for s in sheets]
        if len(labels)!=len(set(labels)) or labels.count(sheet)!=1:raise ValueError('Exact unique source sheet required')
        selected=next(s for s in sheets if s.get('name')==sheet)
        rid=selected.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
        rels=ET.fromstring(archive.read('xl/_rels/workbook.xml.rels')).findall('r:Relationship',NS)
        matches=[r for r in rels if r.get('Id')==rid]
        if len(matches)!=1 or matches[0].get('TargetMode')=='External':raise ValueError('Exact internal sheet relationship required')
        target=matches[0].get('Target','')
        target=target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/'+target)
        if not target.startswith('xl/worksheets/') or '..' in target.split('/'):raise ValueError('Contained worksheet required')
        xml=ET.fromstring(archive.read(target));physical=[];cells={}
        for row in xml.findall('s:sheetData/s:row',NS):
            raw_row=row.get('r','')
            if not re.fullmatch('[1-9][0-9]*',raw_row):raise ValueError('Explicit physical row coordinates required')
            number=int(raw_row)
            if physical and number<=physical[-1]:raise ValueError('Duplicate/unordered worksheet rows')
            physical.append(number)
            for cell in row.findall('s:c',NS):
                coordinate=cell.get('r','')
                if (not re.fullmatch('[A-Z]+'+raw_row,coordinate) or coordinate in cells):raise ValueError('Duplicate/conflicting physical cell coordinate')
                cells[coordinate]=dict(xml_value=cell.findtext('s:v',None,NS),
                    formula=cell.findtext('s:f',None,NS),has_formula=cell.find('s:f',NS) is not None)
    if header_row not in physical:raise ValueError('Physical header row required')
    book=openpyxl.load_workbook(BytesIO(data),read_only=True,data_only=False,keep_links=False)
    try:
        ws=book[sheet];ws.reset_dimensions()
        headers=None;column=None;descriptors=[]
        for row_number,row in enumerate(ws.iter_rows(),1):
            if row_number==header_row:
                headers=[c.value for c in row]
                if headers.count(id_column)!=1:raise ValueError('Unique historical application identifier header required')
                column=headers.index(id_column)
                continue
            if row_number<=header_row or row_number not in physical:continue
            if headers is None:raise ValueError('Header precedes all selected data rows')
            cell=row[column] if column<len(row) else None
            value=cell.value if cell else None;ctype=cell.data_type if cell else 'n';fmt=cell.number_format if cell else 'General'
            coordinate=get_column_letter(column+1)+str(row_number);source=cells.get(coordinate,{})
            key=exact_identifier(value,cell_type=ctype,number_format=fmt,numeric_policy=numeric_policy)
            reason=None
            if source.get('has_formula'):reason='formula_cached_value_not_identity';key=None
            elif ctype=='e':reason='error_cell_not_identity';key=None
            elif value is None or value=='':reason='blank_identifier';key=None
            elif key is None:reason='unsupported_identifier_type_format_or_text'
            observed=value if type(value) in (str,int,float,bool,type(None)) else repr(value)
            descriptors.append(dict(row=row_number,coordinate=coordinate,value=observed,value_python_type=type(value).__name__,
                cell_type=ctype,number_format=fmt,raw_xml_value=source.get('xml_value'),
                formula=source.get('formula'),cached_value_raw=source.get('xml_value') if source.get('has_formula') else None,
                identifier=key,diagnostic=reason))
    finally:book.close()
    if headers is None or len(descriptors)!=sum(r>header_row for r in physical):raise ValueError('Incomplete physical worksheet extraction')
    # Header unknowns are preserved by value/type rather than made into field names.
    header_values=[dict(value=v if type(v) in (str,int,float,bool,type(None)) else repr(v),value_python_type=type(v).__name__) for v in headers]
    extraction=dict(headers=header_values,id_column_letter=get_column_letter(column+1),rows=descriptors)
    keys=[r['identifier'] for r in descriptors if r['identifier'] is not None]
    if len(keys)!=len(set(keys)):raise ValueError('Ambiguous historical application identifier; duplicate source rows require reconciliation')
    return dict(extraction=extraction,extraction_sha256=fingerprint(extraction),data_row_count=len(descriptors),
                identifier_count=len(keys),diagnostic_count=len(descriptors)-len(keys),keys_sha256=fingerprint(sorted(keys)))


def project_packet(folder,trusted_upstream_sha256,*,expected_manifest_sha256,authority_binding,
                   authority_binding_sha256,publications,generation,workbook_scope,workbook_scope_sha256):
    """Both trusted upstream routes are supported; no self-pinned source authority."""
    _sha(trusted_upstream_sha256);_sha(expected_manifest_sha256)
    authority=authority_contract(authority_binding,authority_binding_sha256)
    root=Path(folder)
    if root.is_symlink():raise ValueError('Symlink capture root refused')
    root=root.resolve();mp=root/'manifest.json'
    if mp.is_symlink():raise ValueError('Symlink manifest refused')
    manifest_bytes=mp.read_bytes()
    if hashlib.sha256(manifest_bytes).hexdigest()!=expected_manifest_sha256:raise ValueError('Independent manifest pin mismatch')
    m=strict_json_loads(manifest_bytes);governing_envelope(m)
    if m.get('kind')!='historical_nhmrc_crossref_retention_v1' or m.get('namespace')!='nhmrc' or m.get('complete') is not True or m.get('next_links')!=[]:
        raise ValueError('Complete historical NHMRC packet required')
    exact_text(m.get('batch'));files=m.get('files')
    if not isinstance(files,list):raise ValueError('Complete file inventory required')
    _count(m.get('file_count'),len(files));captured={};entries={}
    for entry in files:
        if not isinstance(entry,dict):raise ValueError('Structured file entry required')
        name=exact_text(entry.get('name'))
        if Path(name).name!=name or '/' in name or '\\' in name or name in ('.','..','manifest.json') or name in entries:raise ValueError('Duplicate/unsafe source filename')
        p=root/name
        if p.is_symlink() or not p.is_file() or p.resolve().parent!=root:raise ValueError('Contained regular source file required')
        raw=p.read_bytes();_count(entry.get('bytes'),len(raw))
        if hashlib.sha256(raw).hexdigest()!=_sha(entry.get('sha256')):raise ValueError('Retained source hash or length drift')
        entries[name]=entry;captured[name]=raw
    if {p.name for p in root.iterdir()}!=set(captured)|{'manifest.json'}:raise ValueError('Complete capture inventory required')
    used=set()
    def referenced(ref):
        if not isinstance(ref,dict):raise ValueError('Structured source reference required')
        name=ref.get('name')
        if not isinstance(name,str) or name not in entries or name in used or ref.get('sha256')!=entries[name]['sha256']:raise ValueError('Missing/repeated or inconsistent source reference')
        if 'bytes' in ref:_count(ref['bytes'],len(captured[name]))
        used.add(name);return captured[name]
    raw_authority=strict_json_loads(referenced(m.get('authority')))
    governing_envelope(raw_authority)
    if (not isinstance(raw_authority,dict) or any(k in raw_authority for k in ('error','errors'))
            or raw_authority.get('id')!=authority['ror'] or m['authority']['sha256']!=authority['authority_raw_sha256']
            or authority['namespace']!='nhmrc'):
        raise ValueError('Independent NHMRC namespace authority binding required')
    links=raw_authority.get('links')
    if not isinstance(links,list) or not any(isinstance(r,dict) and r.get('type')=='website' and r.get('value')=='https://www.nhmrc.gov.au/' for r in links):
        raise ValueError('Official NHMRC authority source required')
    funders=[]
    external=raw_authority.get('external_ids')
    if not isinstance(external,list) or any(not isinstance(r,dict) for r in external):raise ValueError('Structured authority external IDs required')
    for record in external:
        if record.get('type')=='fundref':
            if not isinstance(record.get('all'),list):raise ValueError('Explicit FundRef array required')
            for value in record['all']:funders.append(_doi('10.13039/'+exact_text(value)))
    if len(funders)!=len(set(funders)) or set(funders)!=set(authority['fundref_dois']):raise ValueError('Independent FundRef/ROR binding mismatch')
    source_bytes=referenced(m.get('source_manifest'));upstream=strict_json_loads(source_bytes);governing_envelope(upstream)
    if (upstream.get('format')!='rims-funding-native-workbook-commit-v1' or upstream.get('source')!='nhmrc'
            or upstream.get('complete') is not True or upstream.get('fidelity')!='original_binary_content_exact_saved_readback'):
        raise ValueError('Retained official NHMRC workbook source contract required')
    spec=m.get('workbook')
    if not isinstance(spec,dict):raise ValueError('Exact workbook scope required')
    from urllib.parse import urlparse
    url=urlparse(exact_text(spec.get('source_url')))
    from urllib.parse import unquote
    if any(part in ('.','..') for part in unquote(url.path).split('/')):raise ValueError('Exact official workbook path required')
    if url.scheme!='https' or url.netloc!='www.nhmrc.gov.au' or not url.path.startswith('/sites/default/files/') or url.query or url.fragment or url.params:
        raise ValueError('Official workbook source URL required')
    source_files=upstream.get('files')
    if not isinstance(source_files,list) or any(not isinstance(r,dict) for r in source_files):raise ValueError('Complete upstream file inventory required')
    if 'file_count' in upstream:_count(upstream['file_count'],len(source_files))
    names=[exact_text(r.get('name')) for r in source_files]
    if len(names)!=len(set(names)):raise ValueError('Duplicate upstream file identity')
    selected=[r for r in source_files if r['name']==spec.get('name')]
    if len(selected)!=1 or selected[0].get('source_url')!=spec['source_url']:raise ValueError('Upstream workbook identity drift')
    for entry in source_files:
        if 'sha256' in entry:_sha(entry['sha256'])
        if 'bytes' in entry:_count(entry['bytes'],entry['bytes'])
    binding=selected[0]
    if ('sha256' in binding)!=('bytes' in binding):raise ValueError('Partial upstream workbook hash/length binding')
    route='source_manifest'
    if 'sha256' in binding:
        if m['source_manifest']['sha256']!=trusted_upstream_sha256:raise ValueError('External upstream authority digest mismatch')
        if 'source_retention' in m:raise ValueError('Explicit upstream route required; unused alternate receipt')
    else:
        route='source_retention'
        if not isinstance(m.get('source_retention'),dict) or m['source_retention'].get('sha256')!=trusted_upstream_sha256:raise ValueError('External upstream authority digest mismatch')
        receipt=strict_json_loads(referenced(m['source_retention']));governing_envelope(receipt)
        if receipt.get('status')!='PASS_NATIVE_ORIGINAL_WORKBOOK_ARCHIVE' or not upstream.get('run') or receipt.get('run')!=upstream['run']:raise ValueError('Upstream retention identity drift')
        exact_text(upstream['run']);rf=receipt.get('files')
        if not isinstance(rf,list) or any(not isinstance(r,dict) for r in rf):raise ValueError('Complete retention inventory required')
        if 'file_count' in receipt:_count(receipt['file_count'],len(rf))
        for entry in rf:
            _sha(entry.get('sha256'));_count(entry.get('bytes'),entry.get('bytes'))
        rn=[exact_text(r.get('name')) for r in rf]
        if len(rn)!=len(set(rn)):raise ValueError('Duplicate upstream retention file')
        originals=[r for r in rf if r['name']=='manifest.json'];retained=[r for r in rf if r['name']==spec['name']]
        if len(originals)!=1 or len(retained)!=1:raise ValueError('Upstream retention selection drift')
        if originals[0].get('sha256')!=hashlib.sha256(source_bytes).hexdigest():raise ValueError('Upstream retention manifest drift')
        _count(originals[0].get('bytes'),len(source_bytes))
        if retained[0].get('matches_original_CLI_capture') is not True:raise ValueError('Upstream original-capture comparison missing')
        binding=retained[0]
    data=referenced(spec)
    if binding.get('sha256')!=hashlib.sha256(data).hexdigest():raise ValueError('Upstream workbook digest mismatch')
    _count(binding.get('bytes'),len(data))
    scope=copy.deepcopy(workbook_scope)
    if not isinstance(scope,dict) or fingerprint(scope)!=workbook_scope_sha256:raise ValueError('Independent extraction scope pin required')
    governing_envelope(scope)
    if (scope.get('complete') is not True or scope.get('next_links')!=[] or scope.get('workbook_sha256')!=hashlib.sha256(data).hexdigest()
            or scope.get('sheet')!=spec.get('sheet') or scope.get('id_column')!=spec.get('id_column')):raise ValueError('Complete workbook scope mismatch')
    inspection=inspect_workbook(data,sheet=scope['sheet'],header_row=scope.get('header_row'),id_column=scope['id_column'],numeric_policy=scope.get('numeric_id_policy'))
    if scope.get('id_column_letter')!=inspection['extraction']['id_column_letter']:raise ValueError('Identifier column-coordinate drift')
    for field in ('data_row_count','identifier_count','diagnostic_count'):_count(scope.get(field),inspection[field])
    if any(scope.get(k)!=inspection[k] for k in ('extraction_sha256','keys_sha256')):raise ValueError('Complete extraction/key reconciliation drift')
    if not isinstance(publications,list) or not isinstance(generation,dict):raise ValueError('Complete canonical generation required')
    publications,generation=copy.deepcopy((publications,generation));governing_envelope(generation)
    if (generation.get('complete') is not True or generation.get('next_links')!=[] or generation.get('source_snapshot_id')!=m['batch']
            or generation.get('manifest_sha256')!=expected_manifest_sha256 or generation.get('publications_sha256')!=fingerprint(publications)):
        raise ValueError('Unbound canonical/source generation')
    for k in ('generation_id','canonical_snapshot_id'):exact_text(generation.get(k))
    _count(generation.get('publication_count'),len(publications))
    index={};keys=set()
    for p in publications:
        if not isinstance(p,dict):raise ValueError('Structured canonical records required')
        key=exact_text(p.get('CanonicalKey'))
        if key in keys:raise ValueError('Duplicate canonical key')
        keys.add(key)
        if p.get('PMID') not in (None,'') and (not isinstance(p['PMID'],str) or not re.fullmatch('[1-9][0-9]*',p['PMID'])):raise ValueError('Exact PMID text required')
        if p.get('DOI') not in (None,''):index.setdefault(_doi(p['DOI']),set()).add(key)
    queries=m.get('crossref')
    if not isinstance(queries,list):raise ValueError('Complete Crossref query inventory required')
    _count(m.get('crossref_count'),len(queries));_count(generation.get('crossref_count'),len(queries))
    common=dict(ReviewState='Not reviewed',AuthenticationState='Not authenticated',Confidence='candidate',IdentityQualified=False,
        CanonicalAdmission=False,MutationAuthorized=False,NativeG06Admission=False,SourceVersion=expected_manifest_sha256,SourceVersionKind='retained_manifest_sha256',SourceBatch=m['batch'],
        SourceSnapshotID=m['batch'],CanonicalInputSHA256=generation['publications_sha256'],WorkbookRawSHA256=spec['sha256'],RecordKind='historical_nhmrc_application',
        GenerationID=generation['generation_id'],CanonicalSnapshotID=generation['canonical_snapshot_id'],SourceManifestSHA256=expected_manifest_sha256,
        AuthorityBindingSHA256=authority_binding_sha256,WorkbookScopeSHA256=workbook_scope_sha256,WorkbookExtractionSHA256=inspection['extraction_sha256'],
        ProjectionRevisionSHA256=fingerprint([expected_manifest_sha256,authority_binding_sha256,workbook_scope_sha256,generation['publications_sha256']]))
    applications={r['identifier']:r for r in inspection['extraction']['rows'] if r['identifier'] is not None};result={};dois=set()
    for cell in inspection['extraction']['rows']:
        if cell['diagnostic']:
            key='workbook:'+spec['sha256']+'|'+cell['coordinate']
            result[key]=dict(common,EvidenceKey=key,RelationshipKind='historical_workbook_cell_diagnostic',FromKey='',ToKey='',CanonicalKey='',CanonicalMatchCount=0,
                SourceWorkID='workbook:'+spec['sha256'],RawSHA256=spec['sha256'],ObservedUTC=None,Diagnostic=cell['diagnostic'],
                PublicationDOI='',PublicationPMID='',FromLabel='',ToLabel='',ORCID='',ORCIDStatus='Absent',ROR=authority['ror'],
                IdentityStatus='Unresolved source cell',SourceURL=spec['source_url'],RecordKind='raw_historical_workbook_cell',
                ReferenceType=json.dumps(cell,ensure_ascii=False,separators=(',',':')),SourcePaths=[scope['sheet']+'!'+cell['coordinate']])
    for query in queries:
        _complete(query);doi=_doi(query.get('doi'))
        if doi in dois:raise ValueError('Duplicate Crossref publication identity')
        dois.add(doi);qurl=urlparse(exact_text(query.get('url')))
        from urllib.parse import unquote
        if qurl.scheme!='https' or qurl.netloc!='api.crossref.org' or qurl.query or qurl.fragment or not qurl.path.startswith('/works/') or _doi(unquote(qurl.path[7:]))!=doi:raise ValueError('Exact Crossref query URL required')
        doc=strict_json_loads(referenced(query));governing_envelope(doc)
        work=doc.get('message',{});governing_envelope(work)
        if not isinstance(work,dict) or doc.get('status')!='ok' or any(k in part for part in (doc,work) for k in ('error','errors')) or _doi(work.get('DOI'))!=doi:raise ValueError('Unique complete Crossref publication identity required')
        funding=work.get('funder',[])
        if not isinstance(funding,list):raise ValueError('Structured funder array required')
        for fi,funder in enumerate(funding):
            if not isinstance(funder,dict):raise ValueError('Structured funder required')
            if funder.get('DOI') in (None,''):continue
            fid=_doi(funder['DOI'])
            if fid not in funders:continue
            awards=funder.get('award',[])
            if not isinstance(awards,list):raise ValueError('Explicit source award array required')
            for ai,original in enumerate(awards or [None]):
                if original not in (None,''):identifier(original)
                match=applications.get(original);kind='publication_government_award_reference' if match else 'publication_government_award_diagnostic'
                target='nhmrc:'+original if match else '';source='crossref:'+doi;canon=index.get(doi,set())
                provenance=dict(relationship='explicit_crossref_funding_to_historical_application',recordKind='historical_nhmrc_application',
                    funderDOI=fid,originalAwardNumber=original,authorityROR=authority['ror'],authorityRawSHA256=m['authority']['sha256'],
                    workbookSourceURL=spec['source_url'],workbookRawSHA256=spec['sha256'],sourceManifestSHA256=m['source_manifest']['sha256'],
                    upstreamPin=trusted_upstream_sha256,upstreamRoute=route,workbookSheet=spec['sheet'],identifierColumn=spec['id_column'],
                    workbookRow=match['row'] if match else None,workbookCell=match,retentionBatch=m['batch'])
                key='|'.join([source,kind,fid,original or '',target]);path='/message/funder/'+str(fi)+'/award/'+str(ai) if awards else '/message/funder/'+str(fi)+'/award'
                row=dict(common,EvidenceKey=key,RelationshipKind=kind,SourceWorkID=source,PublicationDOI=doi,PublicationPMID='',
                    FromKey='doi:'+doi,ToKey=target,FromLabel=fid,ToLabel=original,ORCID='',ORCIDStatus='Absent',ROR=authority['ror'],
                    IdentityStatus='Source assertion; not qualified',SourceURL=query['url'],ObservedUTC=query['observed_utc'],RawSHA256=query['sha256'],
                    CanonicalKey=next(iter(canon)) if len(canon)==1 else '',CanonicalMatchCount=len(canon),
                    Diagnostic=None if match else 'no_exact_historical_application',ReferenceType=json.dumps(provenance,ensure_ascii=False,separators=(',',':')))
                if key in result:
                    if {k:v for k,v in result[key].items() if k!='SourcePaths'}!=row:raise ValueError('Conflicting historical source assertion')
                    result[key]['SourcePaths'].append(path)
                else:result[key]=dict(row,SourcePaths=[path])
    if used!=set(captured):raise ValueError('Unreferenced retained source inventory')
    return sorted(result.values(),key=lambda r:r['EvidenceKey'])
