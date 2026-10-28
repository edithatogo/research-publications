import hashlib
import json
from pathlib import Path
import sys
import tempfile
import unittest

import openpyxl

from scripts.government_award_crosswalks import project_packet as raw_project_packet,inspect_workbook
from scripts.source_relationship_evidence import fingerprint

_CONTEXT={}
PUBS=[{'DOI':'10.1234/fictional','CanonicalKey':'fixture:publication','ReviewStatus':'Protected human review'}]


def arguments(root):
    m=json.loads((root/'manifest.json').read_bytes());data=(root/m['workbook']['name']).read_bytes()
    inspection=inspect_workbook(data,numeric_policy='positive_integer_decimal')
    scope=dict(complete=True,next_links=[],sheet='GRANTS DATA',id_column='APP ID',header_row=1,
        id_column_letter='A',numeric_id_policy='positive_integer_decimal',workbook_sha256=hashlib.sha256(data).hexdigest(),
        **{k:inspection[k] for k in ('data_row_count','identifier_count','diagnostic_count','keys_sha256','extraction_sha256')})
    manifest_sha=hashlib.sha256((root/'manifest.json').read_bytes()).hexdigest();binding=_CONTEXT[str(root)]
    return dict(expected_manifest_sha256=manifest_sha,authority_binding=binding,authority_binding_sha256=fingerprint(binding),
        workbook_scope=scope,workbook_scope_sha256=fingerprint(scope),publications=PUBS,
        generation=dict(complete=True,next_links=[],source_snapshot_id=m['batch'],generation_id='fixture-generation',
            canonical_snapshot_id='fixture-canonical',manifest_sha256=manifest_sha,publications_sha256=fingerprint(PUBS),
            publication_count=len(PUBS),crossref_count=len(m['crossref'])))


def anchored_project_packet(root,pin):
    return raw_project_packet(root,pin,**arguments(root))


def project_packet(root):
    return anchored_project_packet(root,_TRUSTED[str(root)])


_TRUSTED = {}


class GovernmentCrosswalkTests(unittest.TestCase):
    def fixture(self, folder, duplicate=False):
        root = Path(folder)
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = 'GRANTS DATA'
        sheet.append(['APP ID', 'Grant Title'])
        sheet.append([1234567, 'Fictional historical award'])
        if duplicate:
            sheet.append([1234567, 'Conflicting historical award'])
        workbook.save(root / 'historical.xlsx')
        workbook.close()
        url = 'https://www.nhmrc.gov.au/sites/default/files/fictional.xlsx'
        source = dict(format='rims-funding-native-workbook-commit-v1', source='nhmrc', complete=True,
                      fidelity='original_binary_content_exact_saved_readback',
                      files=[dict(name='historical.xlsx', source_url=url,
                                  bytes=(root / 'historical.xlsx').stat().st_size,
                                  sha256=hashlib.sha256((root / 'historical.xlsx').read_bytes()).hexdigest())])
        authority = dict(id='https://ror.org/fictional', links=[dict(type='website', value='https://www.nhmrc.gov.au/')],
                         external_ids=[dict(type='fundref', all=['501100000925'])])
        crossref = dict(status='ok', message={'DOI': '10.1234/fictional',
                            'funder': [{'DOI': '10.13039/501100000925', 'award': ['1234567']}]})
        for name, value in [('source.json', source), ('authority.json', authority), ('crossref.json', crossref)]:
            (root / name).write_text(json.dumps(value))
        files = [dict(name=p.name, bytes=p.stat().st_size, sha256=hashlib.sha256(p.read_bytes()).hexdigest())
                 for p in sorted(root.iterdir())]
        by_name = {r['name']: r for r in files}
        manifest = dict(kind='historical_nhmrc_crossref_retention_v1', namespace='nhmrc', batch='fixture',
                        complete=True,next_links=[],file_count=len(files),crossref_count=1, files=files, authority=by_name['authority.json'],
                        source_manifest=by_name['source.json'],
                        workbook=dict(by_name['historical.xlsx'], source_url=url, sheet='GRANTS DATA', id_column='APP ID'),
                        crossref=[dict(by_name['crossref.json'], doi='10.1234/fictional',
                                       url='https://api.crossref.org/works/10.1234/fictional', observed_utc='2026-01-01T00:00:00Z',complete=True,next_links=[],http_status=200)])
        (root / 'manifest.json').write_text(json.dumps(manifest))
        _TRUSTED[str(root)] = by_name['source.json']['sha256']
        _CONTEXT[str(root)]=dict(complete=True,ror=authority['id'],namespace='nhmrc',fundref_dois=['10.13039/501100000925'],
            authority_raw_sha256=by_name['authority.json']['sha256'],evidence_ref='fixture://independent-authority')
        return root

    def replace(self, root, filename, value):
        (root / filename).write_text(json.dumps(value))
        manifest = json.loads((root / 'manifest.json').read_text())
        data = (root / filename).read_bytes()
        def visit(v):
            if isinstance(v, dict):
                if v.get('name') == filename:
                    v.update(bytes=len(data), sha256=hashlib.sha256(data).hexdigest())
                for x in v.values():
                    visit(x)
            elif isinstance(v, list):
                for x in v:
                    visit(x)
        visit(manifest)
        (root / 'manifest.json').write_text(json.dumps(manifest))

    def test_exact_funding_number_retains_coordinate_and_separate_historical_semantics(self):
        with tempfile.TemporaryDirectory() as d:
            rows = project_packet(self.fixture(d))
            self.assertEqual(len(rows), 1)
            row = rows[0]
            self.assertEqual(row['ToKey'], 'nhmrc:1234567')
            self.assertEqual(row['RelationshipKind'], 'publication_government_award_reference')
            self.assertEqual(row['ReviewState'], 'Not reviewed')
            proof = json.loads(row['ReferenceType'])
            self.assertEqual(proof['workbookRow'], 2)
            self.assertEqual(proof['identifierColumn'], 'APP ID')
            self.assertNotIn('purl.org', json.dumps(rows))

    def test_prefixes_and_other_funder_namespaces_are_not_joined(self):
        with tempfile.TemporaryDirectory() as d:
            root = self.fixture(d)
            doc = json.loads((root / 'crossref.json').read_bytes())
            doc['message']['funder'] = [dict(DOI='10.13039/501100000925', award=['APP1234567', 'GNT1234567']),
                                         dict(DOI='10.13039/999999', award=['1234567'])]
            self.replace(root, 'crossref.json', doc)
            rows=project_packet(root)
            self.assertFalse(any(r['RelationshipKind']=='publication_government_award_reference' for r in rows))
            self.assertTrue(all(not r['NativeG06Admission'] for r in rows))

    def test_duplicate_government_identifier_is_refused(self):
        with tempfile.TemporaryDirectory() as d:
            with self.assertRaises(ValueError):
                project_packet(self.fixture(d, duplicate=True))

    def test_source_hash_and_official_source_manifest_are_required(self):
        with tempfile.TemporaryDirectory() as d:
            root = self.fixture(d)
            data = (root / 'historical.xlsx').read_bytes()
            (root / 'historical.xlsx').write_bytes(data + b'drift')
            with self.assertRaises(ValueError):
                project_packet(root)
            (root / 'historical.xlsx').write_bytes(data)
            source = json.loads((root / 'source.json').read_text())
            source['source'] = 'mrff'
            self.replace(root, 'source.json', source)
            with self.assertRaises(ValueError):
                project_packet(root)

    def test_crossref_work_identity_and_independent_namespace_authority_are_required(self):
        with tempfile.TemporaryDirectory() as d:
            root = self.fixture(d)
            doc = json.loads((root / 'crossref.json').read_text())
            doc['message']['DOI'] = '10.1234/other'
            self.replace(root, 'crossref.json', doc)
            with self.assertRaises(ValueError):
                project_packet(root)


    def test_recomputed_outer_manifest_cannot_rebind_different_upstream_workbook(self):
        for field, value in [('sha256', '0' * 64), ('bytes', 1), ('sha256', None), ('bytes', None)]:
            with self.subTest(field=field, value=value), tempfile.TemporaryDirectory() as d:
                root = self.fixture(d)
                source = json.loads((root / 'source.json').read_text())
                if value is None:
                    source['files'][0].pop(field)
                else:
                    source['files'][0][field] = value
                self.replace(root, 'source.json', source)
                with self.assertRaises(ValueError):
                    project_packet(root)

    def test_changed_workbook_with_consistent_outer_hash_still_refused(self):
        with tempfile.TemporaryDirectory() as d:
            root = self.fixture(d)
            data = (root / 'historical.xlsx').read_bytes() + b'changed bytes'
            (root / 'historical.xlsx').write_bytes(data)
            manifest = json.loads((root / 'manifest.json').read_text())
            for item in [manifest['workbook']] + manifest['files']:
                if item['name'] == 'historical.xlsx':
                    item.update(bytes=len(data), sha256=hashlib.sha256(data).hexdigest())
            (root / 'manifest.json').write_text(json.dumps(manifest))
            with self.assertRaises(ValueError):
                project_packet(root)

    def retention_fixture(self, root):
        source = json.loads((root / 'source.json').read_text())
        entry = source['files'][0]
        retained = dict(entry, matches_original_CLI_capture=True)
        entry.pop('sha256')
        entry.pop('bytes')
        source['run'] = 'fictional-run'
        self.replace(root, 'source.json', source)
        data = (root / 'source.json').read_bytes()
        receipt = dict(status='PASS_NATIVE_ORIGINAL_WORKBOOK_ARCHIVE', run='fictional-run',
                       files=[dict(name='manifest.json', bytes=len(data), sha256=hashlib.sha256(data).hexdigest()), retained])
        raw = json.dumps(receipt).encode()
        (root / 'retention.json').write_bytes(raw)
        manifest = json.loads((root / 'manifest.json').read_text())
        reference = dict(name='retention.json', bytes=len(raw), sha256=hashlib.sha256(raw).hexdigest())
        manifest['files'].append(reference)
        manifest['source_retention'] = reference
        manifest['file_count']=len(manifest['files'])
        (root / 'manifest.json').write_text(json.dumps(manifest))
        _TRUSTED[str(root)] = reference['sha256']
        return receipt

    def test_original_hashless_manifest_can_use_independent_matching_retention_receipt(self):
        with tempfile.TemporaryDirectory() as d:
            root = self.fixture(d)
            self.retention_fixture(root)
            self.assertEqual(len(project_packet(root)), 1)

    def test_retention_receipt_must_bind_original_manifest_run_and_workbook(self):
        for mutation in ['run', 'manifest', 'workbook', 'capture', 'duplicate']:
            with self.subTest(mutation=mutation), tempfile.TemporaryDirectory() as d:
                root = self.fixture(d)
                receipt = self.retention_fixture(root)
                if mutation == 'run': receipt['run'] = 'wrong-run'
                elif mutation == 'manifest': receipt['files'][0]['sha256'] = '0' * 64
                elif mutation == 'workbook': receipt['files'][1]['sha256'] = '0' * 64
                elif mutation == 'capture': receipt['files'][1]['matches_original_CLI_capture'] = False
                else: receipt['files'].append(dict(receipt['files'][1]))
                self.replace(root, 'retention.json', receipt)
                with self.assertRaises(ValueError):
                    project_packet(root)

    def test_replacing_receipt_workbook_and_outer_hashes_cannot_replace_external_pin(self):
        with tempfile.TemporaryDirectory() as d:
            root = self.fixture(d)
            receipt = self.retention_fixture(root)
            pin = _TRUSTED[str(root)]
            workbook = openpyxl.load_workbook(root / 'historical.xlsx')
            workbook['GRANTS DATA']['A2'] = 7654321
            workbook.save(root / 'historical.xlsx')
            workbook.close()
            raw = (root / 'historical.xlsx').read_bytes()
            receipt['files'][1].update(bytes=len(raw), sha256=hashlib.sha256(raw).hexdigest())
            self.replace(root, 'retention.json', receipt)
            manifest = json.loads((root / 'manifest.json').read_text())
            for item in [manifest['workbook']] + manifest['files']:
                if item['name'] == 'historical.xlsx':
                    item.update(bytes=len(raw), sha256=hashlib.sha256(raw).hexdigest())
            (root / 'manifest.json').write_text(json.dumps(manifest))
            with self.assertRaises(ValueError):
                anchored_project_packet(root, pin)

    def test_missing_or_malformed_external_pin_is_refused(self):
        with tempfile.TemporaryDirectory() as d:
            root = self.fixture(d)
            for pin in [None, '', 'x' * 64, True]:
                with self.subTest(pin=pin), self.assertRaises(ValueError):
                    anchored_project_packet(root, pin)

    def commit_fictional_workbook(self,root,data):
        # Explicit fixture-owner handover of a new pinned upstream generation.
        (root/'historical.xlsx').write_bytes(data);sha=hashlib.sha256(data).hexdigest()
        source=json.loads((root/'source.json').read_bytes());source['files'][0].update(bytes=len(data),sha256=sha)
        self.replace(root,'source.json',source)
        manifest=json.loads((root/'manifest.json').read_bytes())
        for entry in [manifest['workbook'],*manifest['files']]:
            if entry['name']=='historical.xlsx':entry.update(bytes=len(data),sha256=sha)
        (root/'manifest.json').write_text(json.dumps(manifest));_TRUSTED[str(root)]=manifest['source_manifest']['sha256']

    def test_independent_manifest_authority_scope_and_generation_pins(self):
        import copy
        for mode in ('manifest','authority','authority-rebind','scope','row-count','keyset','column','generation','generation-error','canonical-drift','canonical-duplicate'):
            with self.subTest(mode=mode),tempfile.TemporaryDirectory() as d:
                root=self.fixture(d);args=arguments(root)
                if mode=='manifest':args['expected_manifest_sha256']='0'*64
                elif mode=='authority':args['authority_binding_sha256']='0'*64
                elif mode=='authority-rebind':
                    a=json.loads((root/'authority.json').read_bytes());a['id']='https://ror.org/other';self.replace(root,'authority.json',a)
                    args['expected_manifest_sha256']=hashlib.sha256((root/'manifest.json').read_bytes()).hexdigest();args['generation']['manifest_sha256']=args['expected_manifest_sha256']
                elif mode in ('scope','row-count','keyset','column'):
                    args['workbook_scope']=copy.deepcopy(args['workbook_scope'])
                    if mode=='scope':args['workbook_scope_sha256']='0'*64
                    else:
                        key={'row-count':'data_row_count','keyset':'keys_sha256','column':'id_column_letter'}[mode]
                        args['workbook_scope'][key]={'row-count':2,'keyset':'0'*64,'column':'B'}[mode]
                        args['workbook_scope_sha256']=fingerprint(args['workbook_scope'])
                elif mode=='generation':args['generation']['publications_sha256']='0'*64
                elif mode=='generation-error':args['generation']['errors']=[]
                elif mode=='canonical-drift':args['publications']=[dict(PUBS[0],ReviewStatus='changed')]
                else:
                    args['publications']=[dict(PUBS[0]),dict(PUBS[0])];args['generation'].update(publication_count=2,publications_sha256=fingerprint(args['publications']))
                with self.assertRaises(ValueError):raw_project_packet(root,_TRUSTED[str(root)],**args)

    def test_formula_cache_error_blanks_and_ambiguous_numeric_format_are_visible(self):
        from io import BytesIO
        from zipfile import ZipFile,ZIP_DEFLATED
        import xml.etree.ElementTree as ET
        with tempfile.TemporaryDirectory() as d:
            root=self.fixture(d);w=openpyxl.load_workbook(root/'historical.xlsx');s=w['GRANTS DATA']
            s.append(['=1234567','Formula']);s.append(['#VALUE!','Error']);s.append([12,'Formatted']);s['A5'].number_format='0000000';s.append([None,'Unknown']);s.append(['0123456','Leading zeros'])
            b=BytesIO();w.save(b);w.close()
            # Preserve an actual cached value while proving it is never an identifier.
            original=ZipFile(BytesIO(b.getvalue()));result=BytesIO()
            with ZipFile(result,'w',ZIP_DEFLATED) as out:
                for name in original.namelist():
                    data=original.read(name)
                    if name=='xl/worksheets/sheet1.xml':
                        doc=ET.fromstring(data);ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                        formula=next(c for c in doc.findall('.//s:c',ns) if c.get('r')=='A3');formula.find('s:v',ns).text='1234567';data=ET.tostring(doc)
                    out.writestr(name,data)
            original.close();self.commit_fictional_workbook(root,result.getvalue())
            rows=project_packet(root);diags=[r for r in rows if r['RelationshipKind']=='historical_workbook_cell_diagnostic']
            self.assertEqual(len(diags),4)
            cells={json.loads(r['ReferenceType'])['coordinate']:json.loads(r['ReferenceType']) for r in diags}
            self.assertEqual(cells['A3']['cached_value_raw'],'1234567');self.assertIsNone(cells['A3']['identifier'])
            self.assertEqual(cells['A4']['cell_type'],'e');self.assertEqual(cells['A5']['number_format'],'0000000');self.assertEqual(cells['A6']['diagnostic'],'blank_identifier')
            inspection=inspect_workbook(result.getvalue(),numeric_policy='positive_integer_decimal')
            self.assertEqual(inspection['identifier_count'],2);self.assertEqual(inspection['data_row_count'],6)
            self.assertIn('0123456',[r['identifier'] for r in inspection['extraction']['rows']])
            self.assertTrue(all(not r['MutationAuthorized'] and not r['NativeG06Admission'] for r in rows))

    def test_numeric_identity_requires_explicit_policy_and_never_guesses_zeros(self):
        from scripts.government_award_crosswalks import exact_identifier
        self.assertIsNone(exact_identifier(1234567,cell_type='n'))
        self.assertEqual(exact_identifier(1234567,cell_type='n',numeric_policy='positive_integer_decimal'),'1234567')
        for value in (True,1234567.0,0,-1):self.assertIsNone(exact_identifier(value,cell_type='n',numeric_policy='positive_integer_decimal'))
        self.assertIsNone(exact_identifier(12,cell_type='n',number_format='0000000',numeric_policy='positive_integer_decimal'))
        self.assertEqual(exact_identifier('0123456'),'0123456');self.assertIsNone(exact_identifier(' 0123456'))
        self.assertEqual(exact_identifier('APP1234567'),'APP1234567')

    def test_declared_dimension_cannot_truncate_physical_extraction(self):
        from io import BytesIO
        from zipfile import ZipFile,ZIP_DEFLATED
        with tempfile.TemporaryDirectory() as d:
            root=self.fixture(d);data=(root/'historical.xlsx').read_bytes();out=BytesIO()
            with ZipFile(BytesIO(data)) as src,ZipFile(out,'w',ZIP_DEFLATED) as dst:
                for name in src.namelist():
                    raw=src.read(name)
                    if name=='xl/worksheets/sheet1.xml':raw=raw.replace(b'ref="A1:B2"',b'ref="A1:A1"')
                    dst.writestr(name,raw)
            inspection=inspect_workbook(out.getvalue(),numeric_policy='positive_integer_decimal')
            self.assertEqual(inspection['data_row_count'],1);self.assertEqual(inspection['identifier_count'],1)
            self.assertEqual(inspection['extraction']['rows'][0]['coordinate'],'A2')

    def test_strict_saved_json_symlinks_and_governing_failures(self):
        for mode in ('json','nan','extra','symlink','complete','query-error','query-next','upstream-error','bool-count'):
            with self.subTest(mode=mode),tempfile.TemporaryDirectory() as d:
                root=self.fixture(d)
                if mode in ('json','nan'):
                    raw=b'{"status":"ok","status":"failed"}' if mode=='json' else b'{"x":NaN}'
                    (root/'crossref.json').write_bytes(raw);m=json.loads((root/'manifest.json').read_bytes())
                    for entry in [*m['files'],*m['crossref']]:
                        if entry['name']=='crossref.json':entry.update(bytes=len(raw),sha256=hashlib.sha256(raw).hexdigest())
                    (root/'manifest.json').write_text(json.dumps(m))
                elif mode=='extra':(root/'extra.json').write_text('{}')
                elif mode=='symlink':
                    data=(root/'crossref.json').read_bytes();(root/'crossref.json').unlink();(root/'other.json').write_bytes(data);(root/'crossref.json').symlink_to(root/'other.json')
                elif mode=='upstream-error':
                    source=json.loads((root/'source.json').read_bytes());source['errors']=[];self.replace(root,'source.json',source);_TRUSTED[str(root)]=hashlib.sha256((root/'source.json').read_bytes()).hexdigest()
                else:
                    m=json.loads((root/'manifest.json').read_bytes())
                    if mode=='complete':m['partial']=True
                    elif mode=='query-error':m['crossref'][0]['error']=None
                    elif mode=='query-next':m['crossref'][0]['next_links']=['more']
                    else:m['crossref_count']=True
                    (root/'manifest.json').write_text(json.dumps(m))
                with self.assertRaises(ValueError):project_packet(root)

    def test_occurrences_revisions_unknown_time_and_bridge_cardinality_preserve_inputs(self):
        import copy
        with tempfile.TemporaryDirectory() as d:
            root=self.fixture(d);c=json.loads((root/'crossref.json').read_bytes());c['message']['funder'][0]['award']*=2;self.replace(root,'crossref.json',c)
            m=json.loads((root/'manifest.json').read_bytes());m['crossref'][0]['observed_utc']=None;(root/'manifest.json').write_text(json.dumps(m))
            args=arguments(root);saved=copy.deepcopy(args);before={p.name:p.read_bytes() for p in root.iterdir()}
            first=raw_project_packet(root,_TRUSTED[str(root)],**args);self.assertEqual(args,saved);self.assertEqual(before,{p.name:p.read_bytes() for p in root.iterdir()})
            self.assertEqual(first,raw_project_packet(root,_TRUSTED[str(root)],**args));self.assertEqual(len(first),1);self.assertEqual(len(first[0]['SourcePaths']),2)
            self.assertIsNone(first[0]['ObservedUTC']);self.assertEqual(len({r['CanonicalKey'] for r in first if r['ToKey']}),1)
            c['message']['title']=['Revised source'];self.replace(root,'crossref.json',c);second=project_packet(root)
            self.assertEqual(first[0]['EvidenceKey'],second[0]['EvidenceKey']);self.assertNotEqual(first[0]['SourceVersion'],second[0]['SourceVersion'])
            from scripts.grants_relationships import normalize_relationship
            with self.assertRaises(ValueError):normalize_relationship(first[0])

    def test_raw_pin_paths_and_partial_upstream_binding_do_not_substitute_routes(self):
        with tempfile.TemporaryDirectory() as d:
            root=self.fixture(d);self.retention_fixture(root)
            source=json.loads((root/'source.json').read_bytes());source['files'][0]['sha256']='0'*64;self.replace(root,'source.json',source)
            with self.assertRaises(ValueError):project_packet(root)

    def test_duplicate_headers_and_unmatched_duplicate_keys_block_complete_scope(self):
        from io import BytesIO
        for mode in ('headers','unmatched-duplicate'):
            with self.subTest(mode=mode),tempfile.TemporaryDirectory() as d:
                root=self.fixture(d);w=openpyxl.load_workbook(root/'historical.xlsx');s=w['GRANTS DATA']
                if mode=='headers':s['B1']='APP ID'
                else:s.append(['other','One']);s.append(['other','Two'])
                b=BytesIO();w.save(b);w.close();self.commit_fictional_workbook(root,b.getvalue())
                with self.assertRaises(ValueError):project_packet(root)

    def test_actual_integer_types_and_referenced_identity_scope_are_required(self):
        import copy
        for mode in ('row-count','header-row','bytes','pmid','requested-doi','returned-doi','funder-doi','query-count'):
            with self.subTest(mode=mode),tempfile.TemporaryDirectory() as d:
                root=self.fixture(d);args=arguments(root)
                if mode in ('row-count','header-row'):
                    args['workbook_scope']=copy.deepcopy(args['workbook_scope']);args['workbook_scope']['data_row_count' if mode=='row-count' else 'header_row']=True
                    args['workbook_scope_sha256']=fingerprint(args['workbook_scope'])
                elif mode=='pmid':
                    args['publications']=[dict(PUBS[0],PMID='13 ')];args['generation']['publications_sha256']=fingerprint(args['publications'])
                elif mode in ('returned-doi','funder-doi'):
                    c=json.loads((root/'crossref.json').read_bytes())
                    if mode=='returned-doi':c['message']['DOI']+=' '
                    else:c['message']['funder'][0]['DOI']+=' '
                    self.replace(root,'crossref.json',c);args=arguments(root)
                else:
                    m=json.loads((root/'manifest.json').read_bytes())
                    if mode=='bytes':m['files'][0]['bytes']=True
                    elif mode=='requested-doi':m['crossref'][0]['doi']=' 10.1234/fictional'
                    else:m['crossref_count']=True
                    (root/'manifest.json').write_text(json.dumps(m));args=arguments(root)
                with self.assertRaises(ValueError):raw_project_packet(root,_TRUSTED[str(root)],**args)

    def test_single_read_buffers_and_explicit_qualified_site_import_closure(self):
        from unittest.mock import patch
        import subprocess,sysconfig
        with tempfile.TemporaryDirectory() as d:
            root=self.fixture(d);args=arguments(root);read=Path.read_bytes;counts={}
            def capture(path):counts[path.name]=counts.get(path.name,0)+1;return read(path)
            with patch.object(Path,'read_bytes',capture):raw_project_packet(root,_TRUSTED[str(root)],**args)
            self.assertEqual(counts,{'manifest.json':1,'authority.json':1,'crossref.json':1,'historical.xlsx':1,'source.json':1})
        repo=Path(__file__).resolve().parents[1];qualified_site=sysconfig.get_path('purelib')
        for cwd,statement in ((repo,'from scripts import government_award_crosswalks'),(repo/'scripts','import government_award_crosswalks')):
            code='import sys;sys.path.append('+repr(qualified_site)+');'+statement
            r=subprocess.run([sys.executable,'-B','-S','-c',code],cwd=cwd,capture_output=True,text=True)
            self.assertEqual(r.returncode,0,r.stderr)
